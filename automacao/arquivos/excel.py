"""Planilhas do Excel (.xlsx), com a openpyxl: sem o Excel instalado na máquina."""

from pathlib import Path

from openpyxl import load_workbook


def _texto(valor) -> str:
    """O valor da célula como texto: o número inteiro sem `.0` (telefone, código) e vazio sem `None`."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def ler_registros(caminho: Path, aba: str | None = None) -> list[dict[str, str]]:
    """
    As linhas da planilha como registros, com o cabeçalho da primeira linha como chave.

    Os títulos perdem os espaços das pontas ("Last Name " vira "Last Name"), colunas sem
    título ficam de fora e as linhas vazias são puladas.
    """
    planilha = load_workbook(caminho, read_only=True, data_only=True)

    try:
        folha = planilha[aba] if aba else planilha.active
        linhas = folha.iter_rows(values_only=True)
        cabecalho = [_texto(titulo) for titulo in next(linhas, [])]

        registros = []
        for linha in linhas:
            registro = {titulo: _texto(valor) for titulo, valor in zip(cabecalho, linha) if titulo}
            if any(registro.values()):
                registros.append(registro)
        return registros
    finally:
        planilha.close()
